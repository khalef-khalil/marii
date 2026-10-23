"""Tiny layout DSL: one spec renders to a native .pptx and to a PIL preview (overflow checks)."""

from __future__ import annotations

import re
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SW, SH = 13.333, 7.5
FONT = "Calibri"

NAVY = (0x14, 0x28, 0x4B)
GREEN = (0x2F, 0x7A, 0x5A)
ORANGE = (0xD9, 0x82, 0x2B)
GRAY = (0x6B, 0x77, 0x88)
INK = (0x1F, 0x29, 0x37)
WHITE = (0xFF, 0xFF, 0xFF)
LIGHT = (0xF3, 0xF5, 0xF8)
LGREEN = (0xE3, 0xF1, 0xEA)
LORANGE = (0xFB, 0xEB, 0xD9)
LNAVY = (0xE6, 0xEC, 0xF5)
LINE = (0xC9, 0xD1, 0xDC)

# ----------------------------------------------------------------------------- spec helpers


def runs(text, size, color=INK, bold=False, italic=False):
    out = []
    for i, part in enumerate(text.split("**")):
        if part == "":
            continue
        out.append(dict(t=part, size=size, color=color, bold=bold or i % 2 == 1, italic=italic))
    return out


def P(text, size=18, color=INK, bold=False, italic=False, align="l", bullet=False, after=8, indent=0.26):
    return dict(runs=runs(text, size, color, bold, italic), align=align, bullet=bullet, after=after, indent=indent)


def Box(x, y, w, h, paras=None, fill=None, line=None, line_w=1.0, shape="rect", anchor="t",
        margins=(0.12, 0.08, 0.12, 0.08), radius=0.10):
    return dict(k="box", x=x, y=y, w=w, h=h, paras=paras or [], fill=fill, line=line, line_w=line_w,
                shape=shape, anchor=anchor, margins=margins, radius=radius)


def Pic(path, x, y, w=None, h=None):
    return dict(k="pic", path=str(path), x=x, y=y, w=w, h=h)


def Arrow(x, y, w, h, direction="right", fill=GRAY):
    return dict(k="arrow", x=x, y=y, w=w, h=h, dir=direction, fill=fill)


def Table(x, y, cols, rows, size=14, min_row=0.4, header=True, zebra=True, header_fill=NAVY):
    return dict(k="table", x=x, y=y, cols=cols, rows=rows, size=size, min_row=min_row,
                header=header, zebra=zebra, header_fill=header_fill)


def Cell(text, bold=False, fill=None, color=None, align="l"):
    return dict(text=text, bold=bold, fill=fill, color=color, align=align)


# ----------------------------------------------------------------------------- font metrics (preview)

PX = 120  # preview pixels per inch
_FONT_DIR = Path("/System/Library/Fonts/Supplemental")
_FILES = {
    (False, False): "Arial.ttf",
    (True, False): "Arial Bold.ttf",
    (False, True): "Arial Italic.ttf",
    (True, True): "Arial Bold Italic.ttf",
}
_cache: dict = {}


def pfont(size_pt, bold, italic):
    key = (round(size_pt * 4), bold, italic)
    if key not in _cache:
        _cache[key] = ImageFont.truetype(str(_FONT_DIR / _FILES[(bold, italic)]), max(1, int(round(size_pt / 72 * PX))))
    return _cache[key]


def tokens(text):
    return re.findall(r"[^ \n]+| +", text)


def layout(paras, width_in):
    """Return (lines, total_height_px). Each line: dict(items, h, align, x0)."""
    width_px = width_in * PX
    lines = []
    total = 0.0
    for pi, para in enumerate(paras):
        ind = para["indent"] * PX if para["bullet"] else 0.0
        avail = width_px - ind
        cur, cur_w, cur_h = [], 0.0, 0.0
        para_lines = []

        def flush(first_flag=False):
            nonlocal cur, cur_w, cur_h
            if cur:
                while cur and cur[-1][0].strip() == "":
                    cur_w -= cur[-1][3]
                    cur.pop()
            para_lines.append(dict(items=cur, w=cur_w, h=cur_h or para["runs"][0]["size"] / 72 * PX * 1.2,
                                   align=para["align"], x0=ind))
            cur, cur_w, cur_h = [], 0.0, 0.0

        for r in para["runs"]:
            f = pfont(r["size"], r["bold"], r["italic"])
            for tok in tokens(r["t"]):
                w = f.getlength(tok)
                if tok.strip() != "" and cur_w + w > avail and cur_w > 0:
                    flush()
                if tok.strip() == "" and cur_w == 0:
                    continue
                cur.append((tok, f, r["color"], w))
                cur_w += w
                cur_h = max(cur_h, r["size"] / 72 * PX * 1.2)
        flush()
        if para["bullet"] and para_lines:
            para_lines[0]["bullet"] = (para["runs"][0]["color"], para["runs"][0]["size"])
        lines.extend(para_lines)
        total += sum(l["h"] for l in para_lines)
        if pi < len(paras) - 1:
            total += para["after"] / 72 * PX
            lines.append(dict(items=[], w=0, h=para["after"] / 72 * PX, align="l", x0=0, spacer=True))
    return lines, total


def draw_lines(draw, lines, x, y, w, h_total, anchor, total_h):
    if anchor == "m":
        y += (h_total - total_h) / 2
    elif anchor == "b":
        y += h_total - total_h
    cy = y
    for ln in lines:
        if ln.get("spacer"):
            cy += ln["h"]
            continue
        if ln["align"] == "c":
            cx = x + ln["x0"] + (w - ln["x0"] - ln["w"]) / 2
        elif ln["align"] == "r":
            cx = x + w - ln["w"]
        else:
            cx = x + ln["x0"]
        if ln.get("bullet"):
            col, sz = ln["bullet"]
            draw.text((x, cy), "•", font=pfont(sz, False, False), fill=col)
        for tok, f, col, tw in ln["items"]:
            draw.text((cx, cy), tok, font=f, fill=col)
            cx += tw
        cy += ln["h"]
    return cy


# ----------------------------------------------------------------------------- table geometry


def cell_paras(cell, size, header_row):
    if isinstance(cell, str):
        cell = Cell(cell)
    color = cell["color"] or (WHITE if header_row else INK)
    bold = cell["bold"] or header_row
    return cell, [P(cell["text"], size, color, bold, align=cell["align"], after=0)]


def table_row_heights(spec):
    heights = []
    for ri, row in enumerate(spec["rows"]):
        hmax = spec["min_row"]
        for ci, c in enumerate(row):
            _, paras = cell_paras(c, spec["size"], spec["header"] and ri == 0)
            _, th = layout(paras, spec["cols"][ci] - 0.20)
            hmax = max(hmax, th / PX + 0.14)
        heights.append(round(hmax, 3))
    return heights


# ----------------------------------------------------------------------------- slide model


class Slide:
    def __init__(self, title=None, notes="", footer=True, section=None):
        self.title, self.notes, self.footer, self.section = title, notes, footer, section
        self.el = []
        self.number = 0
        self.footer_label = None

    def add(self, *els):
        self.el.extend(els)
        return self


def chrome(slide, number, label):
    els = []
    if slide.title:
        els.append(Box(0.6, 0.3, 12.1, 0.85, [P(slide.title, 28, NAVY, True, after=0)], anchor="m", margins=(0, 0, 0, 0)))
        els.append(Box(0.6, 1.2, 1.2, 0.06, fill=GREEN))
    if slide.footer:
        els.append(Box(0.6, 7.05, 9.0, 0.3, [P(label, 10, GRAY, after=0)], anchor="m", margins=(0, 0, 0, 0)))
        els.append(Box(11.9, 7.05, 0.8, 0.3, [P(str(number), 10, GRAY, align="r", after=0)], anchor="m", margins=(0, 0, 0, 0)))
    return els


# ----------------------------------------------------------------------------- pptx rendering


def rgb(c):
    return RGBColor(*c)


def fill_tf(tf, paras):
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[para["align"]]
        p.space_after = Pt(para["after"])
        p.space_before = Pt(0)
        for r in para["runs"]:
            run = p.add_run()
            run.text = r["t"]
            run.font.name = FONT
            run.font.size = Pt(r["size"])
            run.font.bold = r["bold"]
            run.font.italic = r["italic"]
            run.font.color.rgb = rgb(r["color"])
        if para["bullet"]:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(Emu(Inches(para["indent"])))))
            pPr.set("indent", str(-int(Emu(Inches(para["indent"])))))
            first_run = para["runs"][0]
            clr = etree.SubElement(pPr, qn("a:buClr"))
            etree.SubElement(clr, qn("a:srgbClr")).set("val", "%02X%02X%02X" % GREEN)
            etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
            etree.SubElement(pPr, qn("a:buChar")).set("char", "•")


def set_border(cell, color, width_pt=0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.SubElement(tcPr, qn(tag))
        ln.set("w", str(int(Pt(width_pt))))
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", "%02X%02X%02X" % color)


def render_pptx(slides, out_path, label, title_meta):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
    blank = prs.slide_layouts[6]
    for i, sl in enumerate(slides, 1):
        s = prs.slides.add_slide(blank)
        for el in chrome(sl, sl.number, sl.footer_label or label) + sl.el:
            add_el(s, el)
        if sl.notes:
            s.notes_slide.notes_text_frame.text = sl.notes
    prs.core_properties.title = title_meta["title"]
    prs.core_properties.author = title_meta["author"]
    prs.save(out_path)


def add_el(s, el):
    k = el["k"]
    if k == "box":
        if el["shape"] == "none" or (el["fill"] is None and el["line"] is None and el["paras"]):
            shp = s.shapes.add_textbox(Inches(el["x"]), Inches(el["y"]), Inches(el["w"]), Inches(el["h"]))
        else:
            kind = MSO_SHAPE.ROUNDED_RECTANGLE if el["shape"] == "round" else MSO_SHAPE.RECTANGLE
            shp = s.shapes.add_shape(kind, Inches(el["x"]), Inches(el["y"]), Inches(el["w"]), Inches(el["h"]))
            if el["shape"] == "round":
                shp.adjustments[0] = min(0.5, el["radius"] / min(el["w"], el["h"]))
            shp.shadow.inherit = False
            if el["fill"] is not None:
                shp.fill.solid()
                shp.fill.fore_color.rgb = rgb(el["fill"])
            else:
                shp.fill.background()
            if el["line"] is not None:
                shp.line.color.rgb = rgb(el["line"])
                shp.line.width = Pt(el["line_w"])
            else:
                shp.line.fill.background()
        tf = shp.text_frame
        tf.word_wrap = True
        ml, mt, mr, mb = el["margins"]
        tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = Inches(ml), Inches(mt), Inches(mr), Inches(mb)
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[el["anchor"]]
        if el["paras"]:
            fill_tf(tf, el["paras"])
    elif k == "pic":
        w, h = pic_size(el)
        s.shapes.add_picture(el["path"], Inches(el["x"]), Inches(el["y"]), Inches(w), Inches(h))
    elif k == "arrow":
        kind = {"right": MSO_SHAPE.RIGHT_ARROW, "down": MSO_SHAPE.DOWN_ARROW, "up": MSO_SHAPE.UP_ARROW}[el["dir"]]
        shp = s.shapes.add_shape(kind, Inches(el["x"]), Inches(el["y"]), Inches(el["w"]), Inches(el["h"]))
        shp.shadow.inherit = False
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(el["fill"])
        shp.line.fill.background()
    elif k == "table":
        add_table(s, el)


def pic_size(el):
    with Image.open(el["path"]) as im:
        ar = im.width / im.height
    if el["w"] and el["h"]:
        return el["w"], el["h"]
    if el["w"]:
        return el["w"], el["w"] / ar
    return el["h"] * ar, el["h"]


def add_table(s, spec):
    heights = table_row_heights(spec)
    nrows, ncols = len(spec["rows"]), len(spec["cols"])
    gf = s.shapes.add_table(nrows, ncols, Inches(spec["x"]), Inches(spec["y"]),
                            Inches(sum(spec["cols"])), Inches(sum(heights)))
    tbl = gf.table
    tblPr = gf._element.graphic.graphicData.tbl.tblPr
    for a in ("firstRow", "bandRow"):
        tblPr.set(a, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None:
        sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    for ci, w in enumerate(spec["cols"]):
        tbl.columns[ci].width = Inches(w)
    for ri, h in enumerate(heights):
        tbl.rows[ri].height = Inches(h)
    for ri, row in enumerate(spec["rows"]):
        is_h = spec["header"] and ri == 0
        for ci, c in enumerate(row):
            cell = tbl.cell(ri, ci)
            cd, paras = cell_paras(c, spec["size"], is_h)
            cell.margin_left = cell.margin_right = Inches(0.10)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            set_border(cell, LINE)
            fill = spec["header_fill"] if is_h else (cd["fill"] or (LIGHT if spec["zebra"] and ri % 2 == 0 else WHITE))
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(fill)
            fill_tf(cell.text_frame, paras)
            cell.text_frame.word_wrap = True


# ----------------------------------------------------------------------------- preview rendering


def render_preview(slides, out_dir, label):
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    problems = []
    for sl in slides:
        im = Image.new("RGB", (int(SW * PX), int(SH * PX)), "white")
        d = ImageDraw.Draw(im)
        for el in chrome(sl, sl.number, sl.footer_label or label) + sl.el:
            prev_el(im, d, el, sl, problems)
        im.save(out_dir / f"slide_{sl.number:02d}.png")
    return problems


def prev_el(im, d, el, sl, problems):
    k = el["k"]
    if k == "box":
        x, y, w, h = (el[a] * PX for a in "xywh")
        if el["fill"] is not None or el["line"] is not None:
            rr = el["radius"] * PX
            if el["shape"] == "round":
                d.rounded_rectangle([x, y, x + w, y + h], radius=rr, fill=el["fill"], outline=el["line"], width=max(1, int(el["line_w"])))
            else:
                d.rectangle([x, y, x + w, y + h], fill=el["fill"], outline=el["line"], width=max(1, int(el["line_w"])))
        if el["paras"]:
            ml, mt, mr, mb = (m * PX for m in el["margins"])
            iw = (el["w"] - (el["margins"][0] + el["margins"][2]))
            lines, th = layout(el["paras"], iw)
            ih = h - mt - mb
            if th > ih + 2:
                problems.append((sl.number, "overflow", round((th - ih) / PX, 2), el["paras"][0]["runs"][0]["t"][:40]))
                d.rectangle([x, y, x + w, y + h], outline=(255, 0, 0), width=3)
            draw_lines(d, lines, x + ml, y + mt, w - ml - mr, ih, el["anchor"], th)
    elif k == "pic":
        w, h = pic_size(el)
        with Image.open(el["path"]).convert("RGBA") as p:
            p = p.resize((int(w * PX), int(h * PX)))
            im.paste(p, (int(el["x"] * PX), int(el["y"] * PX)), p)
        if el["y"] + h > 7.0 or el["x"] + w > SW:
            problems.append((sl.number, "pic out of bounds", round(el["y"] + h, 2), el["path"]))
    elif k == "arrow":
        x, y, w, h = (el[a] * PX for a in "xywh")
        if el["dir"] == "right":
            pts = [(x, y + h * .3), (x + w * .55, y + h * .3), (x + w * .55, y), (x + w, y + h / 2),
                   (x + w * .55, y + h), (x + w * .55, y + h * .7), (x, y + h * .7)]
        elif el["dir"] == "down":
            pts = [(x + w * .3, y), (x + w * .7, y), (x + w * .7, y + h * .55), (x + w, y + h * .55),
                   (x + w / 2, y + h), (x, y + h * .55), (x + w * .3, y + h * .55)]
        else:
            pts = [(x + w / 2, y), (x + w, y + h * .45), (x + w * .7, y + h * .45), (x + w * .7, y + h),
                   (x + w * .3, y + h), (x + w * .3, y + h * .45), (x, y + h * .45)]
        d.polygon(pts, fill=el["fill"])
    elif k == "table":
        heights = table_row_heights(el)
        y = el["y"] * PX
        bottom = el["y"] + sum(heights)
        if bottom > 6.95:
            problems.append((sl.number, "table too low", round(bottom, 2), ""))
        for ri, row in enumerate(el["rows"]):
            x = el["x"] * PX
            is_h = el["header"] and ri == 0
            rh = heights[ri] * PX
            for ci, c in enumerate(row):
                cw = el["cols"][ci] * PX
                cd, paras = cell_paras(c, el["size"], is_h)
                fill = el["header_fill"] if is_h else (cd["fill"] or (LIGHT if el["zebra"] and ri % 2 == 0 else WHITE))
                d.rectangle([x, y, x + cw, y + rh], fill=fill, outline=LINE, width=1)
                lines, th = layout(paras, el["cols"][ci] - 0.20)
                draw_lines(d, lines, x + 0.10 * PX, y, cw - 0.20 * PX, rh, "m", th)
                x += cw
            y += rh
